from pptx import Presentation

def extract_slide_texts(pptx_path):
    prs = Presentation(pptx_path)
    slides = []
    for slide in prs.slides:
        text = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
        slides.append(" ".join(text))
    return slides
